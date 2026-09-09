from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = r"C:\masjid-cicangkudu\output\docx\Dokumentasi_Sistem_Masjid_Cicangkudu.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def border(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for side in ('top','left','bottom','right'):
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '4')
        node.set(qn('w:color'), color)
        borders.append(node)
    tcPr.append(borders)

def set_font(run, size=None, bold=False, color=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.bold = bold
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor(*color)

def add_text(doc, text, size=10.5, space_after=7):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_font(r, size=size)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=16 if level == 1 else 12, bold=True, color=(18, 61, 47))
    return p

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i,h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c, '0B5C3D')
        border(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: set_font(r, size=9, bold=True, color=(255,255,255))
    for ri,row in enumerate(rows):
        cells = t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text = str(v)
            if ri % 2: shade(cells[i], 'F4F8F5')
            border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs: set_font(r, size=8.5)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(70)
r = p.add_run('Dokumentasi Sistem Digital Masjid Jami Cicangkudu'); set_font(r, size=24, bold=True, color=(12,62,47))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Panduan Alur Sistem, Struktur Kode, dan Penggunaan Admin'); set_font(r, size=13, color=(80,92,88))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(24)
r = p.add_run('Penyaji: [Nama Anda]'); set_font(r, size=11, bold=True, color=(12,62,47))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('September 2026'); set_font(r, size=10, color=(80,92,88))
doc.add_page_break()

heading(doc, 'Ringkasan Sistem')
add_text(doc, 'Sistem Digital Masjid Jami Cicangkudu adalah aplikasi Laravel untuk warga dan pengelola masjid. Warga memperoleh jadwal salat Tasikmalaya, informasi dan kegiatan, donasi, laporan keuangan, notifikasi, dan profil. Admin mengelola data warga, konten, pembayaran, verifikasi donasi, laporan, profil masjid, dan notifikasi.')
heading(doc, 'Aktor dan Hak Akses')
table(doc, ['Aktor', 'Akses Utama'], [
    ['Warga', 'Melihat jadwal, informasi, donasi, laporan keuangan, notifikasi, dan mengubah nama, kata sandi, serta foto profil.'],
    ['Admin', 'Mengelola warga, informasi dan kegiatan, program donasi, QRIS dan rekening, laporan keuangan, profil masjid, serta verifikasi bukti donasi.'],
], [1.2, 5.7])
heading(doc, 'Alur Sistem')
table(doc, ['Tahap', 'Alur'], [
    ['1. Autentikasi', 'Pengguna login. Middleware membatasi halaman admin hanya untuk role admin.'],
    ['2. Informasi', 'Admin membuat informasi atau kegiatan dan memilih status terbit. Konten tampil pada warga dan menghasilkan notifikasi.'],
    ['3. Donasi', 'Warga memilih program, mengunggah bukti transfer, lalu sistem membuat data berstatus pending.'],
    ['4. Verifikasi', 'Admin menerima atau menolak bukti. Donasi diterima dicatat otomatis sebagai pemasukan laporan keuangan.'],
    ['5. Neraca', 'Saldo kas dihitung dari pemasukan dikurangi pengeluaran. Aset kas dan dana bersih ditampilkan pada neraca kas.'],
], [1.0, 5.9])

heading(doc, 'Struktur Kode')
table(doc, ['Bagian', 'File Penting', 'Tanggung Jawab'], [
    ['Routing', 'routes/web.php', 'Mendaftarkan halaman warga, admin, CRUD konten, verifikasi donasi, dan hapus massal.'],
    ['Konten Admin', 'AdminContentController.php', 'CRUD informasi, kegiatan, program donasi, laporan keuangan, serta hapus massal konten.'],
    ['Donasi', 'DonationController.php dan AdminDonationController.php', 'Unggah bukti, pemberitahuan admin, menerima atau menolak donasi.'],
    ['Keuangan', 'LaporanController.php', 'Menghitung pemasukan, pengeluaran, saldo, dan neraca kas.'],
    ['Profil Masjid', 'MasjidProfileController.php', 'Menyimpan foto, biodata, visi, dan misi yang tampil pada warga.'],
    ['Profil Warga', 'WargaProfileController.php', 'Memperbarui nama, kata sandi, dan foto profil warga.'],
], [1.2, 2.2, 3.5])

heading(doc, 'Struktur Data Utama')
table(doc, ['Tabel', 'Isi'], [
    ['users', 'Akun warga dan admin beserta role.'],
    ['warga_profiles', 'Nomor HP, alamat, dan avatar warga.'],
    ['masjid_contents', 'Informasi, kegiatan, program donasi, serta transaksi laporan keuangan.'],
    ['donations', 'Nominal, metode pembayaran, bukti transfer, status verifikasi, dan relasi warga.'],
    ['payment_settings', 'Rekening bank dan gambar QRIS yang tampil pada halaman donasi.'],
    ['masjid_profiles', 'Profil masjid: foto, biodata, visi, dan misi.'],
    ['notifications', 'Notifikasi bukti donasi untuk admin dan notifikasi konten atau status donasi untuk warga.'],
], [1.7, 5.2])

heading(doc, 'Pengelolaan Neraca Kas')
add_text(doc, 'Sistem menggunakan basis kas. Pemasukan menambah kas dan dana bersih. Pengeluaran mengurangi kas dan dana bersih. Neraca menampilkan Aset berupa Kas dan Bank, Kewajiban bernilai nol sampai modul utang ditambahkan, serta Dana Bersih yang sama dengan saldo kas. Prinsip ini membuat jumlah aset selalu seimbang dengan kewajiban ditambah dana bersih.')
table(doc, ['Komponen', 'Rumus'], [
    ['Total Pemasukan', 'Jumlah semua transaksi pemasukan yang telah diterbitkan.'],
    ['Total Pengeluaran', 'Jumlah semua transaksi pengeluaran yang telah diterbitkan.'],
    ['Kas dan Bank', 'Total Pemasukan - Total Pengeluaran.'],
    ['Dana Bersih', 'Kas dan Bank - Kewajiban.'],
], [2.0, 4.9])

heading(doc, 'Panduan Admin')
table(doc, ['Menu', 'Langkah'], [
    ['Informasi dan Kegiatan', 'Tambah data, pilih jenis konten, isi judul, deskripsi, tanggal, gambar, dan status terbit. Gunakan Profil Masjid untuk mengubah kartu biodata warga.'],
    ['Program Donasi', 'Tambah program, atur rekening dan QRIS, lalu periksa bukti transfer yang masuk.'],
    ['Laporan Keuangan', 'Tambah pemasukan atau pengeluaran. Pantau ringkasan dan Neraca Kas. Donasi terverifikasi masuk otomatis sebagai pemasukan.'],
    ['Data Warga', 'Tambah manual, impor Excel, edit, hapus satu data, atau pilih beberapa warga untuk hapus massal.'],
    ['Hapus Massal', 'Centang data yang diperlukan, klik Hapus data terpilih, lalu konfirmasi. Tindakan tidak dapat dibatalkan.'],
], [1.8, 5.1])

heading(doc, 'Pengujian dan Pemeliharaan')
add_text(doc, 'Setelah perubahan kode, jalankan php artisan migrate untuk menerapkan struktur database baru, php artisan storage:link agar gambar dapat ditampilkan, php artisan view:cache untuk memeriksa template Blade, dan php artisan test untuk menjalankan pengujian dasar. Pastikan data produksi memiliki cadangan database sebelum memakai hapus massal.')

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('Sistem Digital Masjid Jami Cicangkudu')
    set_font(run, size=8, color=(100,100,100))

doc.save(OUT)
